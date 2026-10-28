#!/usr/bin/env python3
"""
md2xlsx.py — tabele z pliku Markdown → arkusze Excela (.xlsx).

Pandoc zamienia Markdown na Word (.docx), ale nie zapisuje plików .xlsx.
Ten skrypt przenosi do Excela każdą tabelę Markdown z pliku: jedna tabela
= jeden arkusz (Tabela1, Tabela2, …). Liczby zostają liczbami.

Przykłady:
    Windows (PowerShell):  py md2xlsx.py matryca.md matryca.xlsx
    macOS (Terminal):      python3 md2xlsx.py matryca.md matryca.xlsx

Wymaga jednorazowej instalacji:
    Windows:  py -m pip install openpyxl
    macOS:    pip3 install openpyxl

Ograniczenia:
    - liczby z przecinkiem dziesiętnym (0,21) zostają tekstem,
    - znak | wewnątrz komórki rozbija wiersz,
    - po konwersji sprawdź liczbę wierszy i jedną liczbę w arkuszu.
"""

import re
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("Brak biblioteki openpyxl. Zainstaluj: py -m pip install openpyxl (Windows) "
             "albo pip3 install openpyxl (macOS).")


def wartosc(komorka: str):
    """Liczba całkowita lub dziesiętna (z kropką) zostaje liczbą, reszta tekstem."""
    try:
        return float(komorka) if "." in komorka else int(komorka)
    except ValueError:
        return komorka


def zbierz_tabele(sciezka: str) -> list[list[str]]:
    """Zwróć listę tabel; tabela to lista kolejnych linii zaczynających się od |."""
    tabele, biezaca = [], []
    with open(sciezka, encoding="utf-8") as plik:
        for linia in plik:
            if linia.strip().startswith("|"):
                biezaca.append(linia.strip())
            elif biezaca:
                tabele.append(biezaca)
                biezaca = []
    if biezaca:
        tabele.append(biezaca)
    return tabele


def main() -> None:
    if len(sys.argv) != 3:
        sys.exit("Użycie: md2xlsx.py plik.md plik.xlsx")
    wejscie, wyjscie = sys.argv[1], sys.argv[2]

    tabele = zbierz_tabele(wejscie)
    if not tabele:
        sys.exit(f"W pliku {wejscie} nie znaleziono tabel Markdown.")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nr, tabela in enumerate(tabele, 1):
        ws = wb.create_sheet(f"Tabela{nr}")
        for wiersz in tabela:
            komorki = [k.strip() for k in wiersz.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", k) for k in komorki):
                continue  # linia oddzielająca nagłówek |---|---|
            ws.append([wartosc(k) for k in komorki])
        print(f"Tabela{nr}: {ws.max_row} wierszy (z nagłówkiem)")

    wb.save(wyjscie)
    print(f"Zapisano {len(tabele)} tabel(e) do {wyjscie}")


if __name__ == "__main__":
    main()
